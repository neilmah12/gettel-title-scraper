#!/usr/bin/env python3
"""
Build the rate lookup package: estimated 5-year multifamily mortgage rate on any date.

    python build_rate_lookup.py [--out rate_lookup] [--goc data/goc_5y.csv]

Reads data/*.csv (GoC yields, CMB issue spreads, insured lender band, conventional spread band)
and writes, to the output folder:
    rate_lookup_daily.csv / rate_lookup_monthly.csv / rate_lookup.xlsx
    data_dictionary.csv / RATE_LOOKUP_README.md
The daily table has one row per CALENDAR day (weekends and holidays carry the last observation
forward) so a plain exact-match lookup on a sale date always finds a row.
"""
import argparse
import os
from datetime import date

import pandas as pd

import refi_rates as rr

LAGS = (30, 60, 90)

DICTIONARY = [
    ("date", "Calendar date (lookup key). Weekends and holidays repeat the previous observation.", "date", "-", "-"),
    ("goc_5y_pct", "Government of Canada 5-year benchmark bond yield, last observation on or before the date.", "%", "Bank of Canada V39053 (= BD.CDN.5YR.DQ.YLD)", "Hard data"),
    ("goc_obs_date", "Date of the GoC observation actually used.", "date", "Bank of Canada", "Hard data"),
    ("cmb_spread_bps", "Canada Mortgage Bond 5-year spread over the GoC, held constant from each quarterly issue date until the next issue.", "bps", "Bank of Canada published (2024-03 onward); issue yield minus GoC close (before 2024)", "Hard data, two methods up to ~6 bps apart"),
    ("cmb_issue_date", "Date of the CMB issue whose spread applies.", "date", "Canada Housing Trust issues", "Hard data"),
    ("cmb_spread_method", "boc_published or calc_daily_close.", "text", "-", "-"),
    ("cmb_yield_est_pct", "GoC 5-yr + CMB spread: estimated 5-year CMB yield (not the lender's rate).", "%", "Derived", "Derived"),
    ("ins_lender_low_bps", "Lender spread over CMB for CMHC-insured multifamily, low case (sharpest pricing).", "bps", "CMLS commentary through 2024Q1; assumed after", "See ins_confidence"),
    ("ins_lender_base_bps", "Same, base (typical) case.", "bps", "As above", "See ins_confidence"),
    ("ins_lender_high_bps", "Same, high case.", "bps", "As above", "See ins_confidence"),
    ("ins_all_in_low_pct", "INSURED estimate, low: GoC + CMB spread + lender low.", "%", "Derived", "Estimate"),
    ("ins_all_in_base_pct", "INSURED estimate, base: GoC + CMB spread + lender base. Use this as the headline insured rate.", "%", "Derived", "Estimate"),
    ("ins_all_in_high_pct", "INSURED estimate, high: GoC + CMB spread + lender high.", "%", "Derived", "Estimate"),
    ("ins_confidence", "Confidence in the lender-spread band for that quarter (sourced / partly sourced / estimated / assumed).", "text", "-", "-"),
    ("conv_spread_low_bps", "Conventional (uninsured) multifamily spread over the GoC, low case.", "bps", "CMLS / Intellifi commentary through 2023; estimated after", "See conv_confidence"),
    ("conv_spread_base_bps", "Same, base case.", "bps", "As above", "See conv_confidence"),
    ("conv_spread_high_bps", "Same, high case.", "bps", "As above", "See conv_confidence"),
    ("conv_all_in_low_pct", "CONVENTIONAL estimate, low: GoC + conventional spread low.", "%", "Derived", "Estimate"),
    ("conv_all_in_base_pct", "CONVENTIONAL estimate, base: GoC + conventional spread base. Headline conventional rate.", "%", "Derived", "Estimate"),
    ("conv_all_in_high_pct", "CONVENTIONAL estimate, high.", "%", "Derived", "Estimate"),
    ("conv_confidence", "Confidence in the conventional spread band for that quarter.", "text", "-", "-"),
] + [(f"{k}_all_in_base_lag{n}_pct", f"{'INSURED' if k == 'ins' else 'CONVENTIONAL'} base estimate {n} days BEFORE the date. "
      "Rates are usually locked weeks before a sale closes or a mortgage registers, so test these as sensitivities.",
      "%", "Derived", "Estimate") for k in ("ins", "conv") for n in LAGS]


def build_daily(goc, tables, band):
    ins = rr.build_daily_series(goc, tables)
    conv = rr.build_conventional_series(goc, band)
    cmb, _ = tables
    cal = pd.date_range(goc.index.min(), goc.index.max(), freq="D")
    d = pd.DataFrame(index=cal)
    d.index.name = "date"
    g = goc.reindex(cal)
    d["goc_5y_pct"] = g.ffill()
    d["goc_obs_date"] = pd.Series(goc.index, index=goc.index).reindex(cal).ffill().dt.date
    iss = cmb.set_index("issue_date").sort_index()
    d["cmb_spread_bps"] = iss["spread_bps"].reindex(cal.union(iss.index)).ffill().reindex(cal)
    d["cmb_issue_date"] = pd.Series(iss.index, index=iss.index).reindex(cal.union(iss.index)).ffill().reindex(cal).dt.date
    method = pd.Series(["boc_published" if pd.notna(b) else "calc_daily_close" for b in iss["boc_published_bps"]], index=iss.index)
    d["cmb_spread_method"] = method.reindex(cal.union(iss.index)).ffill().reindex(cal)
    first_cmb = iss.index.min()
    d.loc[d.index < first_cmb, ["cmb_spread_bps", "cmb_issue_date", "cmb_spread_method"]] = None
    d["cmb_yield_est_pct"] = d["goc_5y_pct"] + d["cmb_spread_bps"] / 100

    def put(src, mapping, since):
        sub = src.reindex(cal).ffill()
        for col_in, col_out in mapping.items():
            d[col_out] = sub[col_in]
        d.loc[d.index < since, list(mapping.values())] = None

    put(ins, {"lender_low_bps": "ins_lender_low_bps", "lender_base_bps": "ins_lender_base_bps",
              "lender_high_bps": "ins_lender_high_bps", "all_in_low_pct": "ins_all_in_low_pct",
              "all_in_base_pct": "ins_all_in_base_pct", "all_in_high_pct": "ins_all_in_high_pct",
              "lender_confidence": "ins_confidence"}, first_cmb)
    put(conv, {"lender_low_bps": "conv_spread_low_bps", "lender_base_bps": "conv_spread_base_bps",
               "lender_high_bps": "conv_spread_high_bps", "all_in_low_pct": "conv_all_in_low_pct",
               "all_in_base_pct": "conv_all_in_base_pct", "all_in_high_pct": "conv_all_in_high_pct",
               "lender_confidence": "conv_confidence"}, conv.attrs["first_covered"])
    # all-in rates need a same-day GoC; recompute from the calendar GoC so weekends match the carried yield
    for c in ("low", "base", "high"):
        d[f"ins_all_in_{c}_pct"] = d["goc_5y_pct"] + (d["cmb_spread_bps"] + d[f"ins_lender_{c}_bps"]) / 100
        d[f"conv_all_in_{c}_pct"] = d["goc_5y_pct"] + d[f"conv_spread_{c}_bps"] / 100
    for k in ("ins", "conv"):
        for n in LAGS:
            d[f"{k}_all_in_base_lag{n}_pct"] = d[f"{k}_all_in_base_pct"].shift(n)
    return d.reset_index()


def build_monthly(daily):
    m = daily.copy()
    m["month"] = pd.to_datetime(m["date"]).dt.to_period("M")
    num = [c for c in m.columns if c.endswith(("_pct", "_bps"))]
    out = m.groupby("month")[num].mean().round(4)
    out["month_label"] = [p.strftime("%b %Y") for p in out.index]
    out["days_in_table"] = m.groupby("month").size()
    out.insert(0, "month", [str(p) for p in out.index])
    return out.reset_index(drop=True)


def readme(daily, monthly, goc_path):
    first_ins = daily.dropna(subset=["ins_all_in_base_pct"])["date"].min()
    first_conv = daily.dropna(subset=["conv_all_in_base_pct"])["date"].min()
    last = daily["date"].max()
    ex_date = pd.Timestamp("2023-10-18")
    ex = daily[daily["date"] == ex_date].iloc[0]
    last_row = daily.iloc[-1]
    from openpyxl.utils import get_column_letter as L
    col = {c: L(i + 1) for i, c in enumerate(daily.columns)}
    return f"""# Rate lookup: estimated 5-year multifamily mortgage rate on any date

Built {date.today():%Y-%m-%d} by `build_rate_lookup.py` (repo: gettel-title-scraper). Coverage: GoC from {daily['date'].min():%Y-%m-%d};
insured estimates from {first_ins:%Y-%m-%d}; conventional estimates from {first_conv:%Y-%m-%d}; all through {last:%Y-%m-%d}.

**These are modeled estimates, not quoted or observed rates.** Use them to size what a loan likely cost around a sale or
registration date. Do not quote them to a client as a lender's rate.

## Files

| File | Use |
|---|---|
| `rate_lookup_daily.csv` | One row per calendar day. Primary lookup table. |
| `rate_lookup_monthly.csv` | Month averages of the same columns, keyed `YYYY-MM` and `Mon YYYY`. |
| `rate_lookup.xlsx` | Both tables, the source tables, the data dictionary and this README in one workbook. |
| `data_dictionary.csv` | Every column: meaning, units, source, confidence. |
| `data/*.csv` (repo) | Inputs: `goc_5y.csv`, `cmb_issue_spreads.csv`, `lender_spread_band.csv`, `conventional_spread_band.csv`. |

## How to look up a rate

Headline columns for a sale date: **`ins_all_in_base_pct`** (if the loan was CMHC-insured) and **`conv_all_in_base_pct`**
(if conventional). `*_low_pct` and `*_high_pct` give the range. Percent units (4.625 = 4.625%).

Excel (the Daily sheet is contiguous, so exact match works; the sale date in A2 must be a real date, not text):
```
=XLOOKUP(A2, Daily!{col['date']}:{col['date']}, Daily!{col['ins_all_in_base_pct']}:{col['ins_all_in_base_pct']}, "no data")
=XLOOKUP(A2, Daily!{col['date']}:{col['date']}, Daily!{col['conv_all_in_base_pct']}:{col['conv_all_in_base_pct']}, "no data")
=IFERROR(VLOOKUP(A2, Daily!A:AB, MATCH("ins_all_in_base_pct", Daily!1:1, 0), FALSE), "no data")
```
Column letters on the Daily sheet: date = {col['date']}, goc_5y_pct = {col['goc_5y_pct']}, ins_all_in_low/base/high = {col['ins_all_in_low_pct']}/{col['ins_all_in_base_pct']}/{col['ins_all_in_high_pct']},
conv_all_in_low/base/high = {col['conv_all_in_low_pct']}/{col['conv_all_in_base_pct']}/{col['conv_all_in_high_pct']}.
Python:
```python
import pandas as pd
rates = pd.read_csv("rate_lookup_daily.csv", parse_dates=["date"])
sales["date"] = pd.to_datetime(sales["Sale Date"]).dt.normalize()
sales = sales.merge(rates[["date", "ins_all_in_base_pct", "conv_all_in_base_pct"]], on="date", how="left")
```
Example: a sale on {ex_date:%Y-%m-%d} gives GoC {ex['goc_5y_pct']:.2f}%, CMB spread {ex['cmb_spread_bps']:.1f} bps,
insured {ex['ins_all_in_low_pct']:.3f}% / **{ex['ins_all_in_base_pct']:.3f}%** / {ex['ins_all_in_high_pct']:.3f}% (low / base / high),
conventional {ex['conv_all_in_low_pct']:.3f}% / **{ex['conv_all_in_base_pct']:.3f}%** / {ex['conv_all_in_high_pct']:.3f}%.
Latest row ({last:%Y-%m-%d}): GoC {last_row['goc_5y_pct']:.2f}%, insured base {last_row['ins_all_in_base_pct']:.3f}%, conventional base {last_row['conv_all_in_base_pct']:.3f}%.

A blank means no estimate exists for that date (before the coverage start).

## How each number is built

- **GoC 5-year** (`goc_5y_pct`): Bank of Canada daily benchmark yield V39053. Hard data.
- **CMB spread** (`cmb_spread_bps`): spread of the 5-year Canada Mortgage Bond over the GoC, from the 20 quarterly Canada Housing Trust
  issues since 2021-12-14, held constant until the next issue. From 2024-03 it is the Bank of Canada published initial spread; before that
  it is the issue yield minus the GoC close on the pricing date (verified to reproduce all nine calculated spreads exactly).
  The two methods can differ by up to about 6 bps.
- **Insured all-in** = GoC + CMB spread + lender spread over CMB. CMHC-insured loans are quoted over the CMB, not the GoC.
  The lender spread is a band (low / base / high) by quarter: sourced from CMLS commentary ranges through 2024Q1, estimated for 2022Q4,
  and **assumed (45 / 55 / 70 bps) from 2024Q2**. 2022Q2's base is an index average, a different definition from neighbouring quarters.
  Replace the assumed rows in `data/lender_spread_band.csv` with real quotes and rebuild.
- **Conventional all-in** = GoC + conventional spread (uninsured multifamily is priced over the GoC with a wider spread).
  Bands by quarter come from CMLS and Intellifi commentary for 2021 to 2023 (150 bps at 2021 year-end, 210 to 250 typical at 2022 year-end,
  160 to 210 through 2023) and are **estimates from 2024** (no public Alberta multifamily series exists).

## Confidence by period

| Period | Insured band | Conventional band |
|---|---|---|
| 2021-12 to 2022-Q1 | partly sourced | partly sourced |
| 2022 | sourced / estimated | partly sourced to sourced |
| 2023 | sourced / partly sourced | sourced / partly sourced |
| 2024-Q1 | sourced | estimated |
| 2024-Q2 onward | **assumed** | **estimated** |

The `ins_confidence` and `conv_confidence` columns carry this row by row. The GoC and CMB spread inputs are hard data in every period;
the uncertainty sits almost entirely in the lender and conventional spreads, which move slowly. Before-and-after comparisons
(for example a loan's original vs renewal rate) are therefore more reliable than absolute levels, because a constant spread cancels.

## Limits and cautions

1. **Five-year terms only.** Ten-year insured loans price over the 10-year CMB and GoC; the CMB spread was roughly 45 bps for ten years vs
   about 30 for five in 2023. Not modeled here.
2. **Sale date is not the rate-lock date.** Rates are usually locked weeks before closing or registration. The `*_lag30/60/90_pct` columns give the
   base estimate 30, 60 and 90 days earlier; use them to see how much a lock-date difference would move the answer.
3. **No coverage before {first_ins:%Y-%m-%d}** (insured) or {first_conv:%Y-%m-%d} (conventional). Sales earlier than that return blanks.
   Extending back needs CMB issue spreads and lender or conventional spread commentary for those years.
4. **Insured vs conventional is not known from the sale record.** Use `ins_*` for loans believed to be CMHC-insured (typically 75%+ leverage,
   40-year amortization, specialized lenders), `conv_*` otherwise, and show both when unsure. The mortgage workflow scores this per loan.
5. **Not an offered rate.** Actual pricing depends on leverage, property quality, borrower, term, lender appetite and insurance product (MLI Select, etc.).
6. **Updating.** Download a fresh Bank of Canada "Selected Bond Yields" file into `data/goc_5y.csv`, add each new CMB issue to `data/cmb_issue_spreads.csv`,
   and run `python build_rate_lookup.py`. GoC rows after the last observation do not exist, so the table ends on the last GoC date.

## Provenance

GoC yields: Bank of Canada V39053 (file `{goc_path[0]}`, last observation {pd.Timestamp(goc_path[1]).date()}).
CMB spreads and the insured lender band: prior analysis (`refi_distress_screen.py`); CMB 2024+ spreads checked against the Bank of Canada CMB_PURCH dataset.
Conventional bands: deep-research summary of CMLS Commercial Mortgage Commentary and Intellifi reports (Q4 2021, Q4 2022, Q1 and Q3 2023); later years estimated.
"""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="rate_lookup")
    ap.add_argument("--goc", default=os.path.join(rr.DATA_DIR, "goc_5y.csv"))
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)

    goc = rr.load_goc(a.goc)
    tables, band = rr.load_tables(), rr.load_conventional_band()
    daily = build_daily(goc, tables, band)
    monthly = build_monthly(daily)
    dd = pd.DataFrame(DICTIONARY, columns=["column", "meaning", "units", "source", "confidence"])
    text = readme(daily, monthly, (os.path.basename(a.goc), goc.index.max()))

    daily.to_csv(os.path.join(a.out, "rate_lookup_daily.csv"), index=False, float_format="%.4f")
    monthly.to_csv(os.path.join(a.out, "rate_lookup_monthly.csv"), index=False, float_format="%.4f")
    dd.to_csv(os.path.join(a.out, "data_dictionary.csv"), index=False)
    with open(os.path.join(a.out, "RATE_LOOKUP_README.md"), "w", encoding="utf-8") as f:
        f.write(text)

    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    cmb, ins_band = tables
    readme_df = pd.DataFrame({"README": text.splitlines()})
    with pd.ExcelWriter(os.path.join(a.out, "rate_lookup.xlsx"), engine="openpyxl") as xw:
        readme_df.to_excel(xw, sheet_name="README", index=False)
        daily.to_excel(xw, sheet_name="Daily", index=False)
        monthly.to_excel(xw, sheet_name="Monthly", index=False)
        dd.to_excel(xw, sheet_name="Data Dictionary", index=False)
        cmb.assign(issue_date=cmb["issue_date"].dt.date).to_excel(xw, sheet_name="CMB Spreads", index=False)
        ins_band.assign(quarter=ins_band["quarter"].astype(str)).to_excel(xw, sheet_name="Insured Lender Band", index=False)
        band.assign(quarter=band["quarter"].astype(str)).to_excel(xw, sheet_name="Conv Spread Band", index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.fill = PatternFill("solid", start_color="1F3864")
                c.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if ws.title == "README":
                ws.column_dimensions["A"].width = 150
                continue
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                head = str(col[0].value or "")
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(len(head) + 2, 12), 60 if ws.title == "Data Dictionary" else 24)
                for c in col[1:]:
                    if head in ("date", "goc_obs_date", "cmb_issue_date", "issue_date"):
                        c.number_format = "yyyy-mm-dd"
                    elif head.endswith("_pct"):
                        c.number_format = "0.000"
                    elif head.endswith("_bps"):
                        c.number_format = "0.0"
    print(f"Wrote {a.out}/ : {len(daily)} daily rows ({daily['date'].min():%Y-%m-%d} to {daily['date'].max():%Y-%m-%d}), {len(monthly)} months")


if __name__ == "__main__":
    main()
