#!/usr/bin/env python3
# %% [markdown]
# # Gettel Workflow: organize -> parse titles -> refi flags -> one Excel file
#
# One file, two ways to use it:
#   * Terminal:  python gettel_workflow.py --base ~/Gettel --db property_database.xlsx
#   * Colab:     paste each `# %%` block into its own cell, edit CONFIG, run top to bottom
#
# Folder layout under CONFIG["base"]:
#   (loose downloads + scraper CSV)   <- inbox, defaults to base
#   Property Sheets/   {PID}.pdf
#   Title Documents/   {PID}_{filename}.pdf
#   Needs Review/      duplicates and conflicts (nothing is ever deleted)
#   _state/            manifest.csv, mortgages.csv, discharges.csv (what has been parsed)

# %%
# Colab only: uncomment, run once
# !pip install -q pdfplumber openpyxl
# from google.colab import drive; drive.mount('/content/drive')

import argparse
import filecmp
import glob
import math
import os
import re
import shutil
import sys
import traceback
import difflib
from datetime import date, datetime

import pandas as pd

try:
    import pdfplumber
except ImportError:  # fallback backend
    pdfplumber = None
    from pypdf import PdfReader

# %%
# ====================== CONFIG: edit per run ======================
CONFIG = {
    "base": "/content/drive/MyDrive/Gettel",       # folder holding PDFs, the CSV and the database
    "inbox": None,                                  # where new downloads land; None = base
    "scraper_csv": None,                            # None = every gettel-titles-*.csv in the inbox
    "database": "property_database.xlsx",           # must contain a "Prop ID" column
    "output": "property_database_with_mortgages.xlsx",
    "cmb_table": None,                              # xlsx/csv with Maturity Month, Estimated 5-Yr CMB Rate,
                                                    # Risk Tier for Refi. None = Sheet2 of the database
    "prop_id_col": "Prop ID",
    "term_years": 5,                                # mortgage term used for the refi window
    "as_of": None,                                  # "YYYY-MM-DD" or None = today
    "reparse": False,                               # True = re-read every title, not just new ones
    "only_parsed": False,                           # True = output only rows that have a title
    "skip_organize": False,
}
# ==================================================================

PROPERTY_DIR = "Property Sheets"
TITLE_DIR = "Title Documents"
REVIEW_DIR = "Needs Review"
STATE_DIR = "_state"

# %% [markdown]
# ## 1. Organize downloads

# %%
DUP_SUFFIX_RE = re.compile(r"\s*\(\d+\)(?=\.pdf$)", re.I)   # browser copies: "name (1).pdf"
PROPERTY_NAME_RE = re.compile(r"^property (\d+)\.pdf$", re.I)


def canon(name):
    """Lower-case name with the browser's ' (1)' duplicate suffix removed."""
    return DUP_SUFFIX_RE.sub("", name).lower()


def resolve(base, path):
    if path is None:
        return None
    return path if os.path.isabs(path) else os.path.join(base, path)


def unique_path(folder, name):
    os.makedirs(folder, exist_ok=True)
    dest = os.path.join(folder, name)
    stem, ext = os.path.splitext(name)
    n = 1
    while os.path.exists(dest):
        dest = os.path.join(folder, f"{stem} ({n}){ext}")
        n += 1
    return dest


def place(group, dest, review_root, log, label):
    """
    Move the newest file in `group` to `dest`. Extra copies, or a file that
    conflicts with an existing different `dest`, go to Needs Review. Returns a
    short result string.
    """
    group = sorted(group, key=os.path.getmtime, reverse=True)
    best, extras = group[0], group[1:]
    result = "moved"
    if os.path.exists(dest):
        if filecmp.cmp(best, dest, shallow=False):
            extras = group                     # identical copy of what we already have
            result = "already organized"
        else:
            shutil.move(best, unique_path(os.path.join(review_root, "conflicts"), os.path.basename(best)))
            log(f"[CONFLICT] {label}: different file already at {os.path.basename(dest)}; new one moved to Needs Review/conflicts")
            result = "conflict"
    else:
        shutil.move(best, dest)
        log(f"[OK]       {os.path.basename(best)} -> {os.path.relpath(dest, os.path.dirname(os.path.dirname(dest)))}")
    for extra in extras:
        if os.path.exists(extra):
            shutil.move(extra, unique_path(os.path.join(review_root, "duplicates"), os.path.basename(extra)))
            log(f"[DUP]      {os.path.basename(extra)} -> Needs Review/duplicates")
    return result


def find_csvs(cfg, inbox):
    if cfg.get("scraper_csv"):
        return [resolve(cfg["base"], cfg["scraper_csv"])]
    return sorted(glob.glob(os.path.join(inbox, "gettel-titles-*.csv")))


def organize(cfg, log=print):
    base = cfg["base"]
    inbox = resolve(base, cfg.get("inbox")) or base
    prop_dir, title_dir, review = (os.path.join(base, d) for d in (PROPERTY_DIR, TITLE_DIR, REVIEW_DIR))
    for d in (prop_dir, title_dir, review):
        os.makedirs(d, exist_ok=True)

    groups = {}
    for f in os.listdir(inbox):
        full = os.path.join(inbox, f)
        if os.path.isfile(full) and f.lower().endswith(".pdf"):
            groups.setdefault(canon(f), []).append(full)

    report = {"moved": 0, "already": 0, "dups_or_conflicts": 0, "missing": [], "no_title": [], "orphans": []}
    handled = set()

    # Property sheets: "Property 47335.pdf" -> Property Sheets/47335.pdf (no CSV row needed)
    for key, files in list(groups.items()):
        m = PROPERTY_NAME_RE.match(key)
        if m:
            res = place(files, os.path.join(prop_dir, f"{m.group(1)}.pdf"), review, log, f"Property {m.group(1)}")
            report["already" if res == "already organized" else "moved"] += 1
            handled.add(key)

    # Titles: the scraper CSV maps PID -> downloaded filename
    csvs = find_csvs(cfg, inbox)
    if not csvs:
        log("[WARN] No scraper CSV found; titles cannot be matched to PIDs.")
    seen_pids = set()
    for csv_path in csvs:
        rows = pd.read_csv(csv_path, dtype=str, encoding="utf-8-sig").fillna("")
        for _, row in rows.iterrows():
            pid, status, fname = row.get("PID", "").strip(), row.get("Status", "").strip(), row.get("Filename", "").strip()
            if not pid or pid in seen_pids:
                continue
            if status.lower() != "downloaded" or not fname:
                report["no_title"].append((pid, status or "blank"))
                continue
            seen_pids.add(pid)
            new_name = fname if fname.startswith(f"{pid}_") else f"{pid}_{fname}"
            dest = os.path.join(title_dir, new_name)
            files = groups.get(canon(fname))
            if files:
                res = place(files, dest, review, log, f"PID {pid}")
                report["already" if res == "already organized" else "moved"] += 1
                handled.add(canon(fname))
            elif not os.path.exists(dest):
                report["missing"].append((pid, fname))
                log(f"[MISS]     {fname} (PID {pid}) not found in inbox")

    report["orphans"] = sorted(os.path.basename(f) for k, fs in groups.items() if k not in handled for f in fs)
    log(f"Organize done. Moved: {report['moved']} | already organized: {report['already']} | "
        f"missing: {len(report['missing'])} | unmatched PDFs left in inbox: {len(report['orphans'])}")
    return report


# %% [markdown]
# ## 2. Parse titles (ENCUMBRANCES, LIENS & INTERESTS)

# %%
ENC_HEADER_RE = re.compile(r"ENCUMBRANCES,?\s*LIENS\s*&\s*INTERESTS", re.I)

# "182 058 105  09/03/2018  MORTGAGE"  or  "232 315 801  18/10/2023  DISCHARGE OF MORTGAGE 222035284"
HEADER_RE = re.compile(
    r"^(?P<regnum>\d{3}\s?\d{3}\s?\d{3})\s+(?P<date>\d{2}/\d{2}/\d{4})\s+"
    r"(?P<etype>[A-Z][A-Z /&,'\-]*?)(?:\s+(?P<ref>\d[\d ]*\d))?\s*$"
)
# Lines repeated on every page or trailing the section; they must not leak into an entry
NOISE_RES = [re.compile(p) for p in (
    r"^-{5,}$", r"^ENCUMBRANCES, LIENS", r"^REGISTRATION(\s*#.*)?$", r"^NUMBER\s+DATE",
    r"^PAGE \d+$", r"^\( CONTINUED \)", r"^# \d{3}", r"^\* ADDITIONAL REGISTRATIONS",
)]
END_RES = [re.compile(p) for p in (r"^TOTAL INSTRUMENTS", r"^THE REGISTRAR OF TITLES CERTIFIES")]

LENDER_RE = re.compile(r"^MORTGAGEE\s*-\s*(.+)", re.I)
LENDER_STOP_RE = re.compile(r"^(\d|P\.?\s?O\.?\s+BOX|ORIGINAL|AGENT|ATTORNEY|MORTGAGEE|CAVEATOR|RE\s*:)", re.I)
AMOUNT_RE = re.compile(r"ORIGINAL\s+PRINCIPAL\s+AMOUNT\s*:?\s*\$?\s*([\d,]+(?:\.\d+)?)", re.I)
REGNUM_RE = re.compile(r"\b(\d{3}\s?\d{3}\s?\d{3})\b")
DISCHARGE_REF_RE = re.compile(r"DISCHARGE\s+OF\s+MORTGAGE\s+(\d{3}\s?\d{3}\s?\d{3})", re.I)
TOTAL_RE = re.compile(r"TOTAL\s+INSTRUMENTS:\s*(\d+)", re.I)
CERTIFIED_RE = re.compile(r"THIS\s+(\d{1,2})\s+DAY\s+OF\s+([A-Z]+),?\s*(\d{4})", re.I)


def extract_pdf_text(path):
    if pdfplumber is not None:
        with pdfplumber.open(path) as pdf:
            return "\n".join((pg.extract_text() or "") for pg in pdf.pages)
    return "\n".join((pg.extract_text() or "") for pg in PdfReader(str(path)).pages)


def norm_regnum(s):
    return re.sub(r"\s+", "", s or "")


def parse_dmy(s):
    try:
        return datetime.strptime(s, "%d/%m/%Y").date()
    except (TypeError, ValueError):
        return None


def parse_certified(text):
    m = CERTIFIED_RE.search(text)
    if not m:
        return None
    try:
        return datetime.strptime(f"{m.group(1)} {m.group(2).title()} {m.group(3)}", "%d %B %Y").date()
    except ValueError:
        return None


def split_entries(section_text):
    """Return [{regnum, date, etype, ref, lines}] for each registered instrument."""
    entries, cur = [], None
    for raw in section_text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if any(r.search(line) for r in END_RES):
            break
        if any(r.search(line) for r in NOISE_RES):
            continue
        m = HEADER_RE.match(line)
        if m:
            if cur:
                entries.append(cur)
            cur = {"regnum": re.sub(r"\s+", " ", m.group("regnum")), "date": m.group("date"),
                   "etype": m.group("etype").strip(), "ref": m.group("ref"), "lines": []}
        elif cur is not None:
            cur["lines"].append(line)
    if cur:
        entries.append(cur)
    return entries


def parse_lenders(lines):
    """MORTGAGEE - NAME. (a name that wraps onto the next line is joined until it ends with '.')"""
    lenders, i = [], 0
    while i < len(lines):
        m = LENDER_RE.match(lines[i])
        if m:
            name, j = m.group(1).strip(), i
            while (not name.endswith(".") and j + 1 < len(lines) and j - i < 2
                   and not LENDER_STOP_RE.match(lines[j + 1])):
                j += 1
                name += " " + lines[j].strip()
            lenders.append(name.rstrip(". ").strip())
            i = j
        i += 1
    return lenders


def parse_mortgage_entry(entry):
    m = AMOUNT_RE.search(" ".join(entry["lines"]))
    return {
        "regnum": entry["regnum"], "regnum_norm": norm_regnum(entry["regnum"]),
        "date": parse_dmy(entry["date"]),
        "lender": "; ".join(parse_lenders(entry["lines"])) or None,
        "amount": float(m.group(1).replace(",", "")) if m else None,
        "discharged": "No", "discharged_by": None, "notes": "",
    }


def parse_discharge_entry(entry):
    block = " ".join([entry["etype"], entry["ref"] or ""] + entry["lines"])
    own = norm_regnum(entry["regnum"])
    ref = None
    if entry["ref"] and len(norm_regnum(entry["ref"])) == 9:
        ref = norm_regnum(entry["ref"])
    if ref is None:
        m = DISCHARGE_REF_RE.search(block) or re.search(r"\bRE\s*:?\s*(\d{3}\s?\d{3}\s?\d{3})", block)
        if m:
            ref = norm_regnum(m.group(1))
    if ref is None:
        for m in REGNUM_RE.finditer(block):
            if norm_regnum(m.group(1)) != own:
                ref = norm_regnum(m.group(1))
                break
    lenders = parse_lenders(entry["lines"])
    return {
        "regnum": entry["regnum"], "date": parse_dmy(entry["date"]), "ref_regnum": ref,
        "ref_lender": lenders[0] if lenders else None,
        "partial": bool(re.search(r"\bPARTIAL\b", block, re.I)),
        "ref_on_title": False, "unmatched": False,
    }


def normalize_name(name):
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9 ]", "", (name or "").upper())).strip()


def names_match(a, b, threshold=0.8):
    na, nb = normalize_name(a), normalize_name(b)
    if not na or not nb:
        return False
    return na == nb or na in nb or nb in na or difflib.SequenceMatcher(None, na, nb).ratio() >= threshold


def match_discharges(mortgages, discharges):
    """
    1. Discharge cites an instrument number that is still listed -> that mortgage is Yes
       (Partial if the discharge is a partial one: the mortgage still encumbers other lands).
    2. Cites a number that is NOT listed -> the mortgage has already dropped off this title.
       Nothing else is changed.
    3. Cites nothing -> fall back to mortgagee name + date; if still unresolved, earlier
       mortgages become Unknown and the discharge is reported for review.
    """
    for d in discharges:
        target = None
        if d["ref_regnum"]:
            target = next((m for m in mortgages if m["regnum_norm"] == d["ref_regnum"]), None)
            if target is None:
                continue
        else:
            cands = [m for m in mortgages if d["ref_lender"] and names_match(m["lender"], d["ref_lender"])
                     and not (d["date"] and m["date"] and m["date"] > d["date"])]
            target = max(cands, key=lambda m: m["date"] or date.min) if cands else None
        if target is not None:
            target["discharged"] = "Partial" if d["partial"] else "Yes"
            target["discharged_by"] = d["regnum"]
            d["ref_on_title"] = True
        elif not d["ref_regnum"]:
            d["unmatched"] = True
            for m in mortgages:
                if m["discharged"] == "No" and not (d["date"] and m["date"] and m["date"] > d["date"]):
                    m["discharged"] = "Unknown"


def process_title(path):
    """Parse one title PDF. Returns dict(pid, mortgages, discharges, certified, warnings)."""
    fname = os.path.basename(path)
    m = re.match(r"^(\d+)_", fname)
    if not m:
        raise ValueError("filename does not start with '{PID}_'")
    text = extract_pdf_text(path)
    warnings = []
    h = ENC_HEADER_RE.search(text)
    entries = split_entries(text[h.end():]) if h else []
    if not h:
        warnings.append("No ENCUMBRANCES section found (scanned or unusual layout?)")
    declared = TOTAL_RE.search(text)
    if declared and int(declared.group(1)) != len(entries):
        warnings.append(f"Title lists {int(declared.group(1))} instruments but {len(entries)} were parsed")

    mortgages = [parse_mortgage_entry(e) for e in entries if e["etype"] == "MORTGAGE"]
    discharges = [parse_discharge_entry(e) for e in entries if "DISCHARGE" in e["etype"] and "MORTGAGE" in e["etype"]]
    match_discharges(mortgages, discharges)
    for mtg in mortgages:
        if mtg["lender"] is None:
            warnings.append(f"Mortgage {mtg['regnum']}: lender not found")
        if mtg["amount"] is None:
            warnings.append(f"Mortgage {mtg['regnum']}: principal amount not found")
    for d in discharges:
        if d["unmatched"]:
            warnings.append(f"Discharge {d['regnum']}: could not tell which mortgage it discharges")
    return {"pid": m.group(1), "file": fname, "mortgages": mortgages, "discharges": discharges,
            "certified": parse_certified(text), "warnings": warnings}


# %% [markdown]
# ## 3. Parse state: only new or changed titles are read

# %%
MANIFEST_COLS = ["pid", "file", "size", "mtime", "status", "parsed_at", "n_mortgages", "n_discharges",
                 "certified", "warnings"]
MORTGAGE_COLS = ["pid", "file", "regnum", "date", "lender", "amount", "discharged", "discharged_by"]
DISCHARGE_COLS = ["pid", "file", "regnum", "date", "ref_regnum", "ref_on_title", "partial", "ref_lender"]


def load_state(base):
    d = os.path.join(base, STATE_DIR)
    os.makedirs(d, exist_ok=True)

    def read(name, cols):
        p = os.path.join(d, name)
        if os.path.exists(p):
            return pd.read_csv(p, dtype={"pid": str, "regnum": str, "ref_regnum": str, "discharged_by": str})
        return pd.DataFrame(columns=cols)
    return read("manifest.csv", MANIFEST_COLS), read("mortgages.csv", MORTGAGE_COLS), read("discharges.csv", DISCHARGE_COLS)


def save_state(base, manifest, mortgages, discharges):
    d = os.path.join(base, STATE_DIR)
    manifest.to_csv(os.path.join(d, "manifest.csv"), index=False)
    mortgages.to_csv(os.path.join(d, "mortgages.csv"), index=False)
    discharges.to_csv(os.path.join(d, "discharges.csv"), index=False)


def parse_titles(cfg, log=print):
    base = cfg["base"]
    manifest, mortgages, discharges = load_state(base)
    title_dir = os.path.join(base, TITLE_DIR)
    files = sorted(f for f in os.listdir(title_dir) if f.lower().endswith(".pdf")) if os.path.isdir(title_dir) else []
    done = {r.file: (r.size, r.mtime) for r in manifest.itertuples() if r.status == "parsed"}
    n_new = n_fail = n_skip = 0

    for f in files:
        path = os.path.join(title_dir, f)
        sig = (os.path.getsize(path), round(os.path.getmtime(path), 2))
        if not cfg.get("reparse") and f in done and (done[f][0], round(done[f][1], 2)) == sig:
            n_skip += 1
            continue
        manifest = manifest[manifest.file != f]
        mortgages = mortgages[mortgages.file != f]
        discharges = discharges[discharges.file != f]
        row = {"file": f, "size": sig[0], "mtime": sig[1], "parsed_at": datetime.now().isoformat(timespec="seconds"),
               "pid": (re.match(r"^(\d+)_", f) or [None, ""])[1]}
        try:
            r = process_title(path)
            manifest = pd.concat([manifest, pd.DataFrame([{**row, "pid": r["pid"], "status": "parsed",
                "n_mortgages": len(r["mortgages"]), "n_discharges": len(r["discharges"]),
                "certified": r["certified"], "warnings": " | ".join(r["warnings"])}])], ignore_index=True)
            mortgages = pd.concat([mortgages, pd.DataFrame(
                [{**{c: m.get(c) for c in MORTGAGE_COLS if c not in ("pid", "file")}, "pid": r["pid"], "file": f}
                 for m in r["mortgages"]], columns=MORTGAGE_COLS)], ignore_index=True)
            discharges = pd.concat([discharges, pd.DataFrame(
                [{**{c: d.get(c) for c in DISCHARGE_COLS if c not in ("pid", "file")}, "pid": r["pid"], "file": f}
                 for d in r["discharges"]], columns=DISCHARGE_COLS)], ignore_index=True)
            n_new += 1
            for w in r["warnings"]:
                log(f"[WARN]  {f}: {w}")
        except Exception as e:  # one bad PDF never stops the batch
            manifest = pd.concat([manifest, pd.DataFrame([{**row, "status": "failed", "n_mortgages": 0, "n_discharges": 0,
                "certified": None, "warnings": f"{type(e).__name__}: {e}"}])], ignore_index=True)
            n_fail += 1
            log(f"[FAIL]  {f}: {type(e).__name__}: {e}")
            traceback.print_exc()

    manifest = manifest[manifest.file.isin(files)]          # forget files that were removed
    mortgages = mortgages[mortgages.file.isin(files)]
    discharges = discharges[discharges.file.isin(files)]
    save_state(base, manifest, mortgages, discharges)
    log(f"Parse done. New/updated: {n_new} | unchanged (skipped): {n_skip} | failed: {n_fail}")
    return manifest, mortgages, discharges


# %% [markdown]
# ## 4. Refi logic

# %%
def add_years(d, years):
    try:
        return d.replace(year=d.year + years)
    except ValueError:                      # 29 Feb
        return d.replace(year=d.year + years, day=28)


def refi_for_pid(mortgages, sale_date, as_of, term_years=5):
    """
    mortgages: list of dicts (date, discharged, regnum, lender ...) for one PID.
    Rule: use the newest mortgage that is not discharged and is still inside its
    term (registered date + term_years is after as_of). A mortgage already past
    its term should have refinanced, so it is flagged and kept off the refi map.
    A mortgage registered before the sale that is still inside its term is kept
    (likely assumed) and noted.
    """
    out = {"include": "No", "flag": "", "mortgage": None, "refi_date": None, "months": None}
    if not mortgages:
        out["flag"] = "No mortgage on title"
        return out
    active = [m for m in mortgages if m["discharged"] != "Yes" and m["date"]]
    if not active:
        out["flag"] = "All mortgages discharged" if all(m["discharged"] == "Yes" for m in mortgages) \
            else "Mortgage date not readable"
        return out

    in_term = [m for m in active if add_years(m["date"], term_years) > as_of]
    if not in_term:
        newest = max(active, key=lambda m: m["date"])
        due = add_years(newest["date"], term_years)
        out["flag"] = (f"Newest active mortgage ({newest['date']:%d/%m/%Y}) is past its {term_years}-yr term "
                       f"(was due {due:%b %Y}); should have refinanced")
        return out

    pick = max(in_term, key=lambda m: m["date"])
    refi = add_years(pick["date"], term_years)
    out.update(include="Yes", mortgage=pick, refi_date=refi,
               months=max(1, math.ceil((refi - as_of).days / 30.4375)))
    notes = []
    if sale_date and pick["date"] < sale_date:
        notes.append("Mortgage pre-dates sale (possible assumption)")
    if pick["discharged"] == "Partial":
        notes.append("Partial discharge registered; mortgage still on title")
    if pick["discharged"] == "Unknown":
        notes.append("A discharge could not be matched; confirm mortgage is still active")
    if len(in_term) > 1:
        notes.append(f"{len(in_term)} active mortgages in term; using newest")
    out["flag"] = "; ".join(notes)
    return out


def refi_window_label(months):
    lo = ((months - 1) // 6) * 6 + 1
    return f"{lo}–{lo + 5} months"


def load_cmb_table(cfg, log=print):
    """Return {'Jul 2028': (rate, tier)} from the CMB table (see CONFIG['cmb_table'])."""
    base = cfg["base"]
    path = resolve(base, cfg.get("cmb_table")) or resolve(base, cfg["database"])
    try:
        if path.lower().endswith(".csv"):
            df = pd.read_csv(path)
        else:
            xl = pd.ExcelFile(path)
            sheet = next((s for s in ("CMB", "Sheet2") if s in xl.sheet_names), xl.sheet_names[0])
            df = xl.parse(sheet)
        key, rate, tier = "Maturity Month", "Estimated 5-Yr CMB Rate", "Risk Tier for Refi"
        df = df.dropna(subset=[key])
        keys = [k.strftime("%b %Y") if isinstance(k, (datetime, pd.Timestamp)) else str(k).strip() for k in df[key]]
        return dict(zip(keys, zip(df[rate], df[tier])))
    except Exception as e:
        log(f"[WARN] CMB table not loaded from {path}: {e}")
        return {}


# %% [markdown]
# ## 5. Merge into the database and write one Excel file

# %%
def build_output(cfg, manifest, mortgages, discharges, org_report=None, log=print):
    base = cfg["base"]
    pid_col = cfg["prop_id_col"]
    as_of = datetime.strptime(cfg["as_of"], "%Y-%m-%d").date() if cfg.get("as_of") else date.today()
    term = int(cfg.get("term_years", 5))
    cmb = load_cmb_table(cfg, log)

    db = pd.read_excel(resolve(base, cfg["database"]), sheet_name=0)
    db[pid_col] = db[pid_col].astype(str).str.strip().str.replace(r"\.0$", "", regex=True)
    sale = pd.to_datetime(db["Sale Date"], errors="coerce") if "Sale Date" in db else pd.Series(pd.NaT, index=db.index)
    db["Sale Date"] = sale
    exceptions = []

    # One title per PID: the most recently certified
    parsed = manifest[manifest.status == "parsed"].copy()
    parsed["certified"] = pd.to_datetime(parsed["certified"], errors="coerce")
    parsed = parsed.sort_values(["pid", "certified", "mtime"], ascending=[True, False, False])
    chosen = parsed.drop_duplicates("pid").set_index("pid")
    for pid, grp in parsed.groupby("pid"):
        if len(grp) > 1:
            exceptions.append(("Multiple titles", pid, ", ".join(grp.file), f"Using {grp.iloc[0].file} (latest certified)"))
    for r in manifest[manifest.status == "failed"].itertuples():
        exceptions.append(("Parse failed", r.pid, r.file, r.warnings))
    for r in chosen.itertuples():
        if isinstance(r.warnings, str) and r.warnings:
            exceptions.append(("Parse warning", r.Index, r.file, r.warnings))
    for pid in sorted(set(chosen.index) - set(db[pid_col])):
        exceptions.append(("Title has no database row", pid, chosen.loc[pid, "file"], ""))
    if org_report:
        for pid, fname in org_report["missing"]:
            exceptions.append(("Title not found in inbox", pid, fname, "Listed as downloaded in the scraper CSV"))
        for pid, status in org_report["no_title"]:
            exceptions.append(("Scraper: no title downloaded", pid, "", f"Status: {status}"))
        for f in org_report["orphans"]:
            exceptions.append(("Unmatched PDF in inbox", "", f, "Not referenced by any scraper CSV"))

    mort_by_pid = {pid: g[g.file == chosen.loc[pid, "file"]] for pid, g in mortgages.groupby("pid") if pid in chosen.index}
    max_n = max([3] + [len(g) for g in mort_by_pid.values()])
    missing_cmb = set()
    rows = []
    for idx, rec in db.iterrows():
        pid, sdate = rec[pid_col], (rec["Sale Date"].date() if pd.notna(rec["Sale Date"]) else None)
        out = {}
        if pid not in chosen.index:
            out.update({"Refi Include": "Pending", "Refi Flag": "No title parsed yet"})
            rows.append(out)
            continue
        info = chosen.loc[pid]
        cert = info.certified.date() if pd.notna(info.certified) else None
        g = mort_by_pid.get(pid, pd.DataFrame(columns=MORTGAGE_COLS))
        mlist = [{**r._asdict(), "date": pd.to_datetime(r.date).date() if pd.notna(r.date) else None}
                 for r in g.itertuples(index=False)]
        mlist.sort(key=lambda m: m["date"] or date.min)
        res = refi_for_pid(mlist, sdate, as_of, term)
        flags = [res["flag"]] if res["flag"] else []
        if cert and sdate and cert < sdate:
            flags.append("Title certified before the sale date; order a newer title")
        out.update({"Refi Include": res["include"], "Refi Flag": "; ".join(flags)})
        if res["refi_date"]:
            key = res["refi_date"].strftime("%b %Y")
            rate, tier = cmb.get(key, (None, None))
            if key not in cmb:
                missing_cmb.add(key)
            out.update({"Refi Window": refi_window_label(res["months"]), "Months to Refi": res["months"],
                        "Est. Refi Date": key, "5 Yr CMB": rate,
                        "Risk": tier if tier is not None else "No CMB data"})
        out.update({"Title File": info.file, "Title Certified": cert})
        for n in range(1, max_n + 1):
            m = mlist[n - 1] if n <= len(mlist) else {}
            out.update({f"Mortgage{n}_Lender": m.get("lender"), f"Mortgage{n}_Date": m.get("date"),
                        f"Mortgage{n}_Amount": m.get("amount"), f"Mortgage{n}_RegNum": m.get("regnum"),
                        f"Mortgage{n}_Discharged": m.get("discharged")})
        rows.append(out)

    new_cols = pd.DataFrame(rows, index=db.index)
    order = ["Refi Include", "Refi Flag", "Refi Window", "Months to Refi", "Est. Refi Date", "5 Yr CMB", "Risk",
             "Title File", "Title Certified"] + [c for c in new_cols.columns if c.startswith("Mortgage")]
    new_cols = new_cols.reindex(columns=order)
    result = pd.concat([db.drop(columns=[c for c in new_cols.columns if c in db.columns]), new_cols], axis=1)
    if cfg.get("only_parsed"):
        result = result[result["Refi Include"] != "Pending"]

    if missing_cmb:
        exceptions.append(("CMB table missing months", "", "", ", ".join(sorted(missing_cmb, key=lambda s: datetime.strptime(s, "%b %Y")))))

    detail = mortgages.merge(chosen.reset_index()[["pid", "file"]], on=["pid", "file"])
    exc_df = pd.DataFrame(exceptions, columns=["Type", "PID", "File", "Detail"])
    write_excel(resolve(base, cfg["output"]), result, detail, discharges, exc_df, cmb)

    n = result["Refi Include"].value_counts().to_dict()
    log(f"Output: {resolve(base, cfg['output'])}")
    log(f"Rows: {len(result)} | Refi Include Yes: {n.get('Yes', 0)} | No: {n.get('No', 0)} | "
        f"Pending (no title): {n.get('Pending', 0)} | Exceptions: {len(exc_df)}")
    return result, exc_df


def write_excel(path, main, detail, discharges, exceptions, cmb):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cmb_df = pd.DataFrame([(k, v[0], v[1]) for k, v in cmb.items()],
                          columns=["Maturity Month", "Estimated 5-Yr CMB Rate", "Risk Tier for Refi"])
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        main.to_excel(xw, sheet_name="Sheet1", index=False)
        detail.to_excel(xw, sheet_name="Mortgage Detail", index=False)
        discharges.to_excel(xw, sheet_name="Discharges", index=False)
        exceptions.to_excel(xw, sheet_name="Exceptions", index=False)
        cmb_df.to_excel(xw, sheet_name="Sheet2", index=False)
        fills = {"Yes": "E2EFDA", "No": "FADADD", "Pending": "F2F2F2"}
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.fill = PatternFill("solid", start_color="1F3864")
                c.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                head = str(col[0].value or "")
                width = min(max([len(str(c.value)) for c in col[:200] if c.value is not None] + [len(head)]) + 2, 45)
                ws.column_dimensions[get_column_letter(col[0].column)].width = max(width, 10)
                for c in col[1:]:
                    if head.endswith("_Date") or head in ("Title Certified", "date"):
                        c.number_format = "dd/mm/yyyy"
                    elif head == "Sale Date":
                        c.number_format = "yyyy-mm-dd"
                    elif head.endswith("_Amount") or head == "amount":
                        c.number_format = "#,##0"
                    elif head in ("5 Yr CMB", "Estimated 5-Yr CMB Rate"):
                        c.number_format = "0.00%"
            if ws.title == "Sheet1":
                ci = [c.value for c in ws[1]].index("Refi Include") + 1
                for row in ws.iter_rows(min_row=2, min_col=ci, max_col=ci):
                    if row[0].value in fills:
                        row[0].fill = PatternFill("solid", start_color=fills[row[0].value])


# %% [markdown]
# ## 6. Run everything

# %%
def run(cfg, log=print):
    cfg = {**CONFIG, **cfg}
    report = None if cfg.get("skip_organize") else organize(cfg, log)
    manifest, mortgages, discharges = parse_titles(cfg, log)
    return build_output(cfg, manifest, mortgages, discharges, report, log)


def _in_notebook():
    return "ipykernel" in sys.modules


def main(argv=None):
    p = argparse.ArgumentParser(description="Organize Gettel PDFs, parse titles, write the refi Excel file.")
    p.add_argument("--base", required=True, help="folder with PDFs, scraper CSV, database")
    p.add_argument("--db", help="property database xlsx (relative to base or absolute)")
    p.add_argument("--output", help="output xlsx name")
    p.add_argument("--csv", help="scraper CSV (default: all gettel-titles-*.csv in the inbox)")
    p.add_argument("--inbox", help="folder where new downloads land (default: base)")
    p.add_argument("--cmb", help="CMB table xlsx/csv (default: Sheet2 of the database)")
    p.add_argument("--as-of", help="YYYY-MM-DD (default: today)")
    p.add_argument("--term-years", type=int)
    p.add_argument("--reparse", action="store_true", help="re-read every title")
    p.add_argument("--only-parsed", action="store_true", help="output only rows that have a title")
    p.add_argument("--skip-organize", action="store_true")
    a = p.parse_args(argv)
    cfg = {"base": a.base, "reparse": a.reparse, "only_parsed": a.only_parsed, "skip_organize": a.skip_organize}
    for k, v in (("database", a.db), ("output", a.output), ("scraper_csv", a.csv), ("inbox", a.inbox),
                 ("cmb_table", a.cmb), ("as_of", a.as_of), ("term_years", a.term_years)):
        if v is not None:
            cfg[k] = v
    run(cfg)


# %%
# Colab: run this cell instead of the CLI
if _in_notebook():
    result, exceptions = run(CONFIG)
elif __name__ == "__main__":
    main()
